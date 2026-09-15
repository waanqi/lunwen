from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

from lxml import etree
from docx import Document


NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "v": "urn:schemas-microsoft-com:vml",
    "o": "urn:schemas-microsoft-com:office:office",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def text_of(node: etree._Element) -> str:
    return "".join(node.xpath(".//w:t/text()", namespaces=NS)).strip()


def main() -> None:
    src = Path(sys.argv[1]).resolve()
    out_json = Path(sys.argv[2]).resolve()
    figures_dir = Path(sys.argv[3]).resolve()
    figures_dir.mkdir(parents=True, exist_ok=True)

    doc = Document(src)
    body_paragraphs = [p.text.strip() for p in doc.paragraphs]
    nonempty_body = [p for p in body_paragraphs if p]
    table_paragraphs = []
    tables = []
    for ti, table in enumerate(doc.tables, 1):
        rows = []
        for row in table.rows:
            values = []
            for cell in row.cells:
                value = "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                values.append(value)
                table_paragraphs.extend(p.text.strip() for p in cell.paragraphs)
            rows.append(values)
        tables.append({"index": ti, "rows": len(table.rows), "columns": max((len(r.cells) for r in table.rows), default=0), "data": rows})

    with zipfile.ZipFile(src) as zf:
        root = etree.fromstring(zf.read("word/document.xml"))
        formulas = root.xpath("//m:oMath", namespaces=NS)
        formula_paragraphs = root.xpath("//m:oMathPara", namespaces=NS)
        ole_nodes = root.xpath("//*[local-name()='OLEObject']")
        drawing_blips = root.xpath("//a:blip", namespaces=NS)
        vml_images = root.xpath("//v:imagedata", namespaces=NS)
        media_names = [n for n in zf.namelist() if n.startswith("word/media/") and not n.endswith("/")]
        embedded_names = [n for n in zf.namelist() if n.startswith("word/embeddings/") and not n.endswith("/")]

        extracted = []
        for name in media_names:
            target = figures_dir / Path(name).name
            target.write_bytes(zf.read(name))
            extracted.append(str(target))

    headings = []
    for p in doc.paragraphs:
        style = p.style.name if p.style else ""
        if p.text.strip() and (style.lower().startswith("heading") or "标题" in style or "title" in style.lower()):
            headings.append({"style": style, "text": p.text.strip()})

    report = {
        "source": str(src),
        "paragraphs": {
            "body_total": len(body_paragraphs),
            "body_nonempty": len(nonempty_body),
            "table_total": len(table_paragraphs),
            "table_nonempty": sum(bool(p) for p in table_paragraphs),
            "all_total": len(body_paragraphs) + len(table_paragraphs),
        },
        "tables": {"count": len(tables), "items": tables},
        "equations": {
            "omath_count": len(formulas),
            "omath_paragraph_count": len(formula_paragraphs),
            "ole_object_count": len(ole_nodes),
            "embedded_object_files": embedded_names,
        },
        "images": {
            "media_file_count": len(media_names),
            "drawing_references": len(drawing_blips),
            "vml_references": len(vml_images),
            "media_files": media_names,
            "extracted_files": extracted,
        },
        "headings": headings,
        "body_text": nonempty_body,
    }
    out_json.parent.mkdir(parents=True, exist_ok=True)
    out_json.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: report[k] for k in ("paragraphs", "tables", "equations", "images", "headings")}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
